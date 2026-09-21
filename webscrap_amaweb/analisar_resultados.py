#!/usr/bin/env python3
"""Consolida resultados AMAWeb e gera tabelas/figuras reproduzíveis.

O script procura todos os JSON sob --entrada, deduplica por URL e conserva a
avaliação mais recente. Bloqueios Cloudflare são resultados da coleta, mas a
nota da página de bloqueio é mantida separada da nota do conteúdo institucional.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlsplit, urlunsplit


STATUS_CONTENT = "Conteúdo institucional avaliado"
STATUS_CLOUDFLARE = "Bloqueio Cloudflare"
STATUS_MISSING = "Sem resultado AMAWeb"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Gera base consolidada, tabelas, estatísticas e figuras dos JSONs AMAWeb."
    )
    parser.add_argument(
        "--entrada",
        type=Path,
        default=Path(__file__).resolve().parent / "saida_mt",
        help="Pasta que contém as rodadas do AMAWeb.",
    )
    parser.add_argument(
        "--catalogo",
        type=Path,
        help="Planilha .xlsx (aba MT_Sites) ou CSV com município, poder e URL.",
    )
    parser.add_argument(
        "--saida",
        type=Path,
        default=Path(__file__).resolve().parent / "analise_atual",
        help="Pasta recriada logicamente a cada execução (arquivos são sobrescritos).",
    )
    parser.add_argument("--top-erros", type=int, default=15, help="Erros exibidos no ranking.")
    parser.add_argument(
        "--sem-graficos",
        action="store_true",
        help="Gera somente bases, tabelas e relatório; não requer matplotlib.",
    )
    return parser.parse_args()


def normalize_url(value: Any) -> str:
    text = str(value or "").strip().rstrip(",")
    if not text:
        return ""
    parts = urlsplit(text)
    if not parts.scheme or not parts.netloc:
        return text.lower().rstrip("/")
    scheme = parts.scheme.lower()
    netloc = parts.netloc.lower()
    if scheme == "https" and netloc.endswith(":443"):
        netloc = netloc[:-4]
    if scheme == "http" and netloc.endswith(":80"):
        netloc = netloc[:-3]
    path = parts.path.rstrip("/")
    return urlunsplit((scheme, netloc, path, parts.query, parts.fragment))


def host_key(value: Any) -> str:
    """Chave tolerante a http/https e www, usada somente quando o domínio é único."""
    host = (urlsplit(str(value or "").strip().rstrip(",")).hostname or "").casefold()
    return host[4:] if host.startswith("www.") else host


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def integer(value: Any) -> int | None:
    parsed = number(value)
    return int(parsed) if parsed is not None else None


def parse_date(value: Any, fallback: float = 0.0) -> tuple[float, str]:
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y %H:%M"):
        try:
            dt = datetime.strptime(text[:19], pattern)
            return dt.timestamp(), dt.isoformat(sep=" ")
        except ValueError:
            pass
    return fallback, text


def is_cloudflare(title: Any) -> bool:
    text = str(title or "").casefold()
    explicit = "cloudflare" in text and any(
        marker in text for marker in ("access denied", "restrict access", "attention required")
    )
    return explicit or "just a moment" in text


def load_xlsx_catalog(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Para ler .xlsx, instale openpyxl: python3 -m pip install openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["MT_Sites"] if "MT_Sites" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    result: list[dict[str, Any]] = []
    for values in rows:
        row = dict(zip(headers, values))
        url = row.get("URL AMAWeb") or row.get("URL") or row.get("url")
        if not url:
            continue
        result.append(
            {
                "ordem": integer(row.get("Ordem")),
                "codigo_municipio": str(row.get("Código Município") or ""),
                "municipio": str(row.get("Município") or ""),
                "poder": str(row.get("Poder") or ""),
                "esfera": str(row.get("Esfera") or ""),
                "unidade_gestora": str(row.get("Unidade Gestora") or ""),
                "indice_transparencia_2025": number(row.get("Índice Transparência 2025")),
                "nivel_transparencia": str(row.get("Nível Transparência") or ""),
                "url": str(url).strip(),
            }
        )
    workbook.close()
    return result


def load_csv_catalog(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(4096)
        stream.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        rows = csv.DictReader(stream, dialect=dialect)
        result = []
        for row in rows:
            lowered = {str(k).strip().casefold(): v for k, v in row.items()}
            url = lowered.get("url amaweb") or lowered.get("url")
            if not url:
                continue
            result.append(
                {
                    "ordem": integer(lowered.get("ordem")),
                    "codigo_municipio": str(lowered.get("código município") or ""),
                    "municipio": str(lowered.get("município") or ""),
                    "poder": str(lowered.get("poder") or ""),
                    "esfera": str(lowered.get("esfera") or ""),
                    "unidade_gestora": str(lowered.get("unidade gestora") or ""),
                    "indice_transparencia_2025": number(
                        lowered.get("índice transparência 2025")
                    ),
                    "nivel_transparencia": str(lowered.get("nível transparência") or ""),
                    "url": str(url).strip(),
                }
            )
    return result


def load_catalog(path: Path | None) -> list[dict[str, Any]]:
    if path is None:
        return []
    if not path.is_file():
        raise RuntimeError(f"Catálogo não encontrado: {path}")
    if path.suffix.casefold() == ".xlsx":
        return load_xlsx_catalog(path)
    return load_csv_catalog(path)


def conform_counts(value: Any) -> tuple[int | None, int | None, int | None]:
    parts = str(value or "").split("@")
    if len(parts) != 3:
        return None, None, None
    try:
        return int(parts[0]), int(parts[1]), int(parts[2])
    except ValueError:
        return None, None, None


def extract_checks(nodes: Any) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    totals = {"testes": 0, "aprovados": 0, "falhos": 0, "manuais": 0}
    failures: list[dict[str, Any]] = []
    if not isinstance(nodes, dict):
        return totals, failures
    for practice, items in nodes.items():
        if not isinstance(items, list):
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            totals["testes"] += 1
            verdict = str(item.get("verdict") or "")
            if verdict == "passed":
                totals["aprovados"] += 1
            elif verdict == "failed":
                totals["falhos"] += 1
                elements = item.get("elements")
                failures.append(
                    {
                        "pratica": str(practice),
                        "codigo": str(item.get("resultCode") or ""),
                        "descricao": str(item.get("description") or ""),
                        "elementos": len(elements) if isinstance(elements, list) else 0,
                    }
                )
            else:
                totals["manuais"] += 1
    return totals, failures


def load_source_urls(root: Path) -> dict[str, str]:
    """Relaciona cada JSON à URL solicitada, inclusive quando houve redirecionamento."""
    sources: dict[str, str] = {}
    for summary in root.rglob("resumo_amaweb*.csv"):
        try:
            with summary.open(encoding="utf-8-sig", newline="") as stream:
                for row in csv.DictReader(stream, delimiter=";"):
                    reference = str(row.get("json") or "").strip()
                    requested_url = str(row.get("url") or "").strip()
                    if not reference or not requested_url:
                        continue
                    json_path = Path(reference)
                    candidates = (
                        [json_path]
                        if json_path.is_absolute()
                        else [summary.parent / json_path, root.parent / json_path, Path.cwd() / json_path]
                    )
                    resolved = next((candidate.resolve() for candidate in candidates if candidate.is_file()), None)
                    if resolved:
                        sources[str(resolved)] = requested_url
        except (OSError, csv.Error):
            continue
    return sources


def load_evaluations(
    root: Path, source_urls: dict[str, str]
) -> tuple[dict[str, dict[str, Any]], list[str], int]:
    if not root.is_dir():
        raise RuntimeError(f"Pasta de entrada não encontrada: {root}")
    latest: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []
    valid_files = 0
    for path in sorted(root.rglob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8-sig"))
            data = payload.get("result", {}).get("data", {})
            if payload.get("success") != 1 or not isinstance(data, dict):
                raise ValueError("estrutura AMAWeb não reconhecida")
            returned_url = str(
                data.get("rawUrl") or data.get("tot", {}).get("info", {}).get("url") or ""
            )
            requested_url = source_urls.get(str(path.resolve()), returned_url)
            key = normalize_url(requested_url)
            if not key:
                raise ValueError("URL ausente")
            timestamp, date_text = parse_date(data.get("date"), path.stat().st_mtime)
            checks, failures = extract_checks(data.get("nodes"))
            err_a, err_aa, err_aaa = conform_counts(data.get("conform"))
            title = str(data.get("title") or "")
            info = data.get("tot", {}).get("info", {})
            candidate = {
                "url": requested_url,
                "url_retornada_amaweb": returned_url,
                "titulo": title,
                "nota_amaweb": number(data.get("score", info.get("score"))),
                "data_avaliacao": date_text,
                "_timestamp": timestamp,
                "resultado_coleta": STATUS_CLOUDFLARE if is_cloudflare(title) else STATUS_CONTENT,
                "erros_a": err_a,
                "erros_aa": err_aa,
                "erros_aaa": err_aaa,
                "elementos_html": integer(info.get("htmlTags")),
                "tamanho_bytes": integer(info.get("size")),
                **checks,
                "falhas_detalhadas": failures,
                "arquivo_json": str(path.resolve()),
                "arquivo_pdf": str(path.with_suffix(".pdf").resolve())
                if path.with_suffix(".pdf").is_file()
                else "",
            }
            valid_files += 1
            previous = latest.get(key)
            if previous is None or candidate["_timestamp"] >= previous["_timestamp"]:
                latest[key] = candidate
        except (OSError, ValueError, TypeError, json.JSONDecodeError) as exc:
            warnings.append(f"{path}: {exc}")
    return latest, warnings, valid_files


def load_attempts(root: Path) -> dict[str, int]:
    attempts: dict[str, int] = {}
    names = {"resumo_amaweb_final.csv", "falhas_amaweb_persistentes.csv"}
    for path in root.rglob("*.csv"):
        if path.name not in names:
            continue
        try:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                for row in csv.DictReader(stream, delimiter=";"):
                    key = normalize_url(row.get("url"))
                    count = integer(row.get("tentativas")) or 0
                    if key:
                        attempts[key] = max(attempts.get(key, 0), count)
        except (OSError, csv.Error):
            continue
    return attempts


def build_records(
    catalog: list[dict[str, Any]], evaluations: dict[str, dict[str, Any]], attempts: dict[str, int]
) -> tuple[list[dict[str, Any]], int]:
    records: list[dict[str, Any]] = []
    catalog_keys: set[str] = set()
    duplicate_catalog = 0
    catalog_by_host: dict[str, list[int]] = defaultdict(list)
    evaluations_by_host: dict[str, list[str]] = defaultdict(list)
    matched: dict[int, tuple[str, str]] = {}
    used_evaluations: set[str] = set()

    for index, item in enumerate(catalog):
        key = normalize_url(item["url"])
        if key in catalog_keys:
            duplicate_catalog += 1
        catalog_keys.add(key)
        catalog_by_host[host_key(item["url"])].append(index)
        if key in evaluations:
            matched[index] = (key, "URL exata")
            used_evaluations.add(key)

    for key, evaluation in evaluations.items():
        if key not in used_evaluations:
            evaluations_by_host[host_key(evaluation["url"])].append(key)

    # Redirecionamentos frequentemente mudam protocolo, www, caminho ou fragmento.
    # O pareamento por domínio só é aceito em relação 1:1 para evitar falsos vínculos.
    for host, indexes in catalog_by_host.items():
        unmatched_indexes = [index for index in indexes if index not in matched]
        candidates = evaluations_by_host.get(host, [])
        if host and len(unmatched_indexes) == 1 and len(candidates) == 1:
            matched[unmatched_indexes[0]] = (candidates[0], "Domínio único após redirecionamento")
            used_evaluations.add(candidates[0])

    for index, item in enumerate(catalog):
        key = normalize_url(item["url"])
        match = matched.get(index)
        evaluation = evaluations.get(match[0]) if match else None
        base = dict(item)
        if evaluation:
            base.update({k: v for k, v in evaluation.items() if not k.startswith("_")})
            base["associacao_url"] = match[1]
        else:
            base.update(
                {
                    "titulo": "",
                    "nota_amaweb": None,
                    "data_avaliacao": "",
                    "resultado_coleta": STATUS_MISSING,
                    "erros_a": None,
                    "erros_aa": None,
                    "erros_aaa": None,
                    "elementos_html": None,
                    "tamanho_bytes": None,
                    "testes": None,
                    "aprovados": None,
                    "falhos": None,
                    "manuais": None,
                    "falhas_detalhadas": [],
                    "arquivo_json": "",
                    "arquivo_pdf": "",
                    "associacao_url": "Sem JSON associado",
                }
            )
        base["tentativas"] = attempts.get(key, 0)
        base["nota_conteudo_institucional"] = (
            base.get("nota_amaweb") if base["resultado_coleta"] == STATUS_CONTENT else None
        )
        records.append(base)

    for key, evaluation in evaluations.items():
        if key in used_evaluations:
            continue
        records.append(
            {
                "ordem": None,
                "codigo_municipio": "",
                "municipio": "",
                "poder": "Não informado",
                "esfera": "",
                "unidade_gestora": "",
                "indice_transparencia_2025": None,
                "nivel_transparencia": "",
                **{k: v for k, v in evaluation.items() if not k.startswith("_")},
                "associacao_url": "JSON fora do catálogo",
                "tentativas": attempts.get(key, 0),
                "nota_conteudo_institucional": evaluation.get("nota_amaweb")
                if evaluation["resultado_coleta"] == STATUS_CONTENT
                else None,
            }
        )
    records.sort(key=lambda row: (row.get("ordem") is None, row.get("ordem") or 10**9, row["url"]))
    return records, duplicate_catalog


def write_csv(path: Path, rows: Iterable[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            clean = {key: "" if row.get(key) is None else row.get(key) for key in fields}
            writer.writerow(clean)


def values(records: Iterable[dict[str, Any]], field: str) -> list[float]:
    return [float(row[field]) for row in records if row.get(field) is not None]


def describe(nums: list[float]) -> dict[str, float | int | None]:
    return {
        "n": len(nums),
        "media": statistics.fmean(nums) if nums else None,
        "mediana": statistics.median(nums) if nums else None,
        "desvio_padrao": statistics.stdev(nums) if len(nums) > 1 else None,
        "minimo": min(nums) if nums else None,
        "maximo": max(nums) if nums else None,
    }


def group_table(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in records:
        groups[row.get("poder") or "Não informado"].append(row)
    output = []
    for name, rows in sorted(groups.items()):
        content_scores = values(rows, "nota_conteudo_institucional")
        all_scores = values(rows, "nota_amaweb")
        desc = describe(content_scores)
        total = len(rows)
        content = sum(row["resultado_coleta"] == STATUS_CONTENT for row in rows)
        cloudflare = sum(row["resultado_coleta"] == STATUS_CLOUDFLARE for row in rows)
        missing = sum(row["resultado_coleta"] == STATUS_MISSING for row in rows)
        output.append(
            {
                "poder": name,
                "total_urls": total,
                "conteudo_avaliado": content,
                "cloudflare": cloudflare,
                "sem_resultado": missing,
                "cobertura_percentual": 100 * (content + cloudflare) / total if total else None,
                "media_conteudo": desc["media"],
                "mediana_conteudo": desc["mediana"],
                "desvio_padrao_conteudo": desc["desvio_padrao"],
                "minimo_conteudo": desc["minimo"],
                "maximo_conteudo": desc["maximo"],
                "media_incluindo_cloudflare": statistics.fmean(all_scores) if all_scores else None,
            }
        )
    return output


def error_table(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    content = [row for row in records if row["resultado_coleta"] == STATUS_CONTENT]
    stats: dict[str, dict[str, Any]] = {}
    for row in content:
        practices_seen: set[str] = set()
        for failure in row.get("falhas_detalhadas", []):
            name = failure["pratica"]
            item = stats.setdefault(
                name,
                {
                    "pratica": name,
                    "codigo_exemplo": failure["codigo"],
                    "descricao_exemplo": failure["descricao"],
                    "sites": 0,
                    "ocorrencias_testes": 0,
                    "elementos_afetados": 0,
                },
            )
            item["ocorrencias_testes"] += 1
            item["elementos_afetados"] += failure["elementos"]
            if name not in practices_seen:
                item["sites"] += 1
                practices_seen.add(name)
    for item in stats.values():
        item["prevalencia_percentual"] = 100 * item["sites"] / len(content) if content else None
    return sorted(stats.values(), key=lambda row: (-row["sites"], row["pratica"]))


def paired_table(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], list[float]] = defaultdict(list)
    for row in records:
        municipality = str(row.get("municipio") or "").strip()
        power = str(row.get("poder") or "")
        score = row.get("nota_conteudo_institucional")
        if municipality and power in {"Executivo", "Legislativo"} and score is not None:
            grouped[(municipality, power)].append(float(score))
    municipalities = sorted({key[0] for key in grouped})
    pairs = []
    for municipality in municipalities:
        executive = grouped.get((municipality, "Executivo"), [])
        legislative = grouped.get((municipality, "Legislativo"), [])
        if not executive or not legislative:
            continue
        score_e = statistics.fmean(executive)
        score_l = statistics.fmean(legislative)
        pairs.append(
            {
                "municipio": municipality,
                "n_executivo": len(executive),
                "nota_executivo": score_e,
                "n_legislativo": len(legislative),
                "nota_legislativo": score_l,
                "diferenca_executivo_menos_legislativo": score_e - score_l,
            }
        )
    return pairs


def rankdata(nums: list[float]) -> list[float]:
    ordered = sorted(enumerate(nums), key=lambda pair: pair[1])
    ranks = [0.0] * len(nums)
    start = 0
    while start < len(ordered):
        end = start + 1
        while end < len(ordered) and ordered[end][1] == ordered[start][1]:
            end += 1
        rank = (start + 1 + end) / 2
        for position in range(start, end):
            ranks[ordered[position][0]] = rank
        start = end
    return ranks


def pearson(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) < 2 or len(xs) != len(ys):
        return None
    mean_x, mean_y = statistics.fmean(xs), statistics.fmean(ys)
    numerator = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
    denominator = math.sqrt(
        sum((x - mean_x) ** 2 for x in xs) * sum((y - mean_y) ** 2 for y in ys)
    )
    return numerator / denominator if denominator else None


def format_number(value: Any, decimals: int = 2) -> str:
    return "—" if value is None else f"{float(value):.{decimals}f}".replace(".", ",")


def latex_escape(value: Any) -> str:
    text = str(value if value is not None else "—")
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "_": r"\_",
        "#": r"\#",
        "{": r"\{",
        "}": r"\}",
    }
    return "".join(replacements.get(char, char) for char in text)


def write_latex(path: Path, headers: list[str], rows: list[list[Any]], alignment: str) -> None:
    lines = [
        "% Gerado automaticamente por analisar_resultados.py",
        f"\\begin{{tabular}}{{{alignment}}}",
        "\\hline",
        " & ".join(latex_escape(value) for value in headers) + r" \\",
        "\\hline",
    ]
    for row in rows:
        lines.append(" & ".join(latex_escape(value) for value in row) + r" \\")
    lines.extend(["\\hline", "\\end{tabular}", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def generate_tables(
    output: Path,
    records: list[dict[str, Any]],
    groups: list[dict[str, Any]],
    errors: list[dict[str, Any]],
    pairs: list[dict[str, Any]],
    valid_json_files: int,
    invalid_json_files: int,
    duplicate_catalog: int,
) -> list[dict[str, Any]]:
    status_counts = Counter(row["resultado_coleta"] for row in records)
    content_scores = values(records, "nota_conteudo_institucional")
    all_scores = values(records, "nota_amaweb")
    general = [
        {"indicador": "URLs no universo analisado", "valor": len(records)},
        {"indicador": "Conteúdos institucionais avaliados", "valor": status_counts[STATUS_CONTENT]},
        {"indicador": "Bloqueios Cloudflare", "valor": status_counts[STATUS_CLOUDFLARE]},
        {"indicador": "Sem resultado AMAWeb", "valor": status_counts[STATUS_MISSING]},
        {"indicador": "Arquivos JSON válidos encontrados", "valor": valid_json_files},
        {"indicador": "Arquivos JSON inválidos", "valor": invalid_json_files},
        {"indicador": "URLs repetidas no catálogo", "valor": duplicate_catalog},
        {"indicador": "Média do conteúdo institucional", "valor": describe(content_scores)["media"]},
        {"indicador": "Mediana do conteúdo institucional", "valor": describe(content_scores)["mediana"]},
        {"indicador": "Média de toda página retornada, inclusive Cloudflare", "valor": describe(all_scores)["media"]},
    ]
    write_csv(output / "tabela_resumo_geral.csv", general, ["indicador", "valor"])
    group_fields = [
        "poder",
        "total_urls",
        "conteudo_avaliado",
        "cloudflare",
        "sem_resultado",
        "cobertura_percentual",
        "media_conteudo",
        "mediana_conteudo",
        "desvio_padrao_conteudo",
        "minimo_conteudo",
        "maximo_conteudo",
        "media_incluindo_cloudflare",
    ]
    write_csv(output / "tabela_por_poder.csv", groups, group_fields)
    write_csv(
        output / "tabela_erros_frequentes.csv",
        errors,
        [
            "pratica",
            "codigo_exemplo",
            "sites",
            "prevalencia_percentual",
            "ocorrencias_testes",
            "elementos_afetados",
            "descricao_exemplo",
        ],
    )
    write_csv(
        output / "tabela_comparacao_executivo_legislativo.csv",
        pairs,
        [
            "municipio",
            "n_executivo",
            "nota_executivo",
            "n_legislativo",
            "nota_legislativo",
            "diferenca_executivo_menos_legislativo",
        ],
    )
    write_latex(
        output / "tabela_por_poder.tex",
        ["Poder", "URLs", "Conteúdo", "Cloudflare", "Sem resultado", "Média", "Mediana"],
        [
            [
                row["poder"],
                row["total_urls"],
                row["conteudo_avaliado"],
                row["cloudflare"],
                row["sem_resultado"],
                format_number(row["media_conteudo"]),
                format_number(row["mediana_conteudo"]),
            ]
            for row in groups
        ],
        "lrrrrrr",
    )
    write_latex(
        output / "tabela_top_erros.tex",
        ["Prática", "Sites", "Prevalência (%)"],
        [
            [row["pratica"], row["sites"], format_number(row["prevalencia_percentual"], 1)]
            for row in errors[:15]
        ],
        "lrr",
    )
    return general


def save_figure(fig: Any, figures: Path, name: str) -> None:
    fig.savefig(figures / f"{name}.png", dpi=300, bbox_inches="tight")
    fig.savefig(figures / f"{name}.pdf", bbox_inches="tight")


def generate_figures(
    output: Path,
    records: list[dict[str, Any]],
    groups: list[dict[str, Any]],
    errors: list[dict[str, Any]],
    pairs: list[dict[str, Any]],
    top_errors: int,
) -> list[str]:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as exc:
        raise RuntimeError(
            "Para gerar figuras, instale matplotlib: python3 -m pip install matplotlib"
        ) from exc

    figures = output / "figuras"
    figures.mkdir(exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "DejaVu Serif",
            "font.size": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.dpi": 120,
        }
    )
    colors = {STATUS_CONTENT: "#0072B2", STATUS_CLOUDFLARE: "#E69F00", STATUS_MISSING: "#7A7A7A"}
    created: list[str] = []

    counts = Counter(row["resultado_coleta"] for row in records)
    labels = [STATUS_CONTENT, STATUS_CLOUDFLARE, STATUS_MISSING]
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    bars = ax.bar(labels, [counts[label] for label in labels], color=[colors[label] for label in labels])
    ax.bar_label(bars, padding=3)
    ax.set_ylabel("Número de URLs")
    ax.set_title("Resultados da coleta automatizada")
    ax.tick_params(axis="x", rotation=12)
    save_figure(fig, figures, "figura_01_resultados_coleta")
    plt.close(fig)
    created.append("figura_01_resultados_coleta")

    content_scores = values(records, "nota_conteudo_institucional")
    if content_scores:
        fig, ax = plt.subplots(figsize=(8.2, 4.4))
        bins = [value / 2 for value in range(0, 22)]
        ax.hist(content_scores, bins=bins, color="#0072B2", edgecolor="white")
        mean = statistics.fmean(content_scores)
        median = statistics.median(content_scores)
        ax.axvline(mean, color="#D55E00", linewidth=2, label=f"Média = {mean:.2f}")
        ax.axvline(median, color="#009E73", linewidth=2, linestyle="--", label=f"Mediana = {median:.2f}")
        ax.set(xlabel="Nota AMAWeb do conteúdo institucional", ylabel="Número de portais", xlim=(0, 10))
        ax.set_title("Distribuição das notas de acessibilidade")
        ax.legend(frameon=False)
        save_figure(fig, figures, "figura_02_distribuicao_notas")
        plt.close(fig)
        created.append("figura_02_distribuicao_notas")

    box_data, box_labels = [], []
    for group in groups:
        scores = [
            float(row["nota_conteudo_institucional"])
            for row in records
            if (row.get("poder") or "Não informado") == group["poder"]
            and row.get("nota_conteudo_institucional") is not None
        ]
        if len(scores) >= 3:
            box_data.append(scores)
            box_labels.append(group["poder"])
    if box_data:
        height = max(4.5, 0.55 * len(box_data) + 1.5)
        fig, ax = plt.subplots(figsize=(8.2, height))
        plot = ax.boxplot(box_data, tick_labels=box_labels, vert=False, patch_artist=True)
        for patch in plot["boxes"]:
            patch.set_facecolor("#56B4E9")
            patch.set_alpha(0.75)
        ax.set(xlabel="Nota AMAWeb do conteúdo institucional", xlim=(0, 10))
        ax.set_title("Notas por poder ou tipo de órgão")
        save_figure(fig, figures, "figura_03_notas_por_poder")
        plt.close(fig)
        created.append("figura_03_notas_por_poder")

    ranked = list(reversed(errors[:top_errors]))
    if ranked:
        fig, ax = plt.subplots(figsize=(8.2, max(5.2, 0.38 * len(ranked) + 1.6)))
        bars = ax.barh(
            [row["pratica"] for row in ranked],
            [row["prevalencia_percentual"] for row in ranked],
            color="#CC79A7",
        )
        ax.bar_label(bars, fmt="%.1f%%", padding=3, fontsize=8)
        ax.set(xlabel="Portais de conteúdo que apresentam o erro (%)", xlim=(0, 105))
        ax.set_title("Erros automáticos mais frequentes")
        save_figure(fig, figures, "figura_04_erros_frequentes")
        plt.close(fig)
        created.append("figura_04_erros_frequentes")

    correlation_rows = [
        row
        for row in records
        if row.get("nota_conteudo_institucional") is not None
        and row.get("indice_transparencia_2025") is not None
    ]
    if len(correlation_rows) >= 3:
        xs = [float(row["indice_transparencia_2025"]) for row in correlation_rows]
        ys = [float(row["nota_conteudo_institucional"]) for row in correlation_rows]
        rho = pearson(rankdata(xs), rankdata(ys))
        fig, ax = plt.subplots(figsize=(7, 5.2))
        ax.scatter(xs, ys, s=24, alpha=0.65, color="#0072B2", edgecolors="none")
        ax.set(xlabel="Índice de Transparência 2025", ylabel="Nota AMAWeb do conteúdo")
        ax.set_title(f"Transparência e acessibilidade (Spearman ρ = {rho:.2f})")
        ax.set_ylim(0, 10.2)
        save_figure(fig, figures, "figura_05_transparencia_acessibilidade")
        plt.close(fig)
        created.append("figura_05_transparencia_acessibilidade")

    if pairs:
        xs = [row["nota_executivo"] for row in pairs]
        ys = [row["nota_legislativo"] for row in pairs]
        fig, ax = plt.subplots(figsize=(6.2, 5.6))
        ax.scatter(xs, ys, s=28, alpha=0.7, color="#009E73", edgecolors="none")
        ax.plot([0, 10], [0, 10], color="#666666", linestyle="--", linewidth=1)
        ax.set(
            xlabel="Nota do Executivo",
            ylabel="Nota do Legislativo",
            xlim=(0, 10.2),
            ylim=(0, 10.2),
        )
        ax.set_title(f"Comparação municipal pareada (n = {len(pairs)})")
        save_figure(fig, figures, "figura_06_executivo_legislativo")
        plt.close(fig)
        created.append("figura_06_executivo_legislativo")
    return created


def generate_report(
    output: Path,
    records: list[dict[str, Any]],
    groups: list[dict[str, Any]],
    errors: list[dict[str, Any]],
    pairs: list[dict[str, Any]],
    warnings: list[str],
    figures: list[str],
    catalog_path: Path | None,
    input_path: Path,
) -> None:
    counts = Counter(row["resultado_coleta"] for row in records)
    content_scores = values(records, "nota_conteudo_institucional")
    all_scores = values(records, "nota_amaweb")
    paired_diffs = [row["diferenca_executivo_menos_legislativo"] for row in pairs]
    corr_rows = [
        row
        for row in records
        if row.get("nota_conteudo_institucional") is not None
        and row.get("indice_transparencia_2025") is not None
    ]
    xs = [float(row["indice_transparencia_2025"]) for row in corr_rows]
    ys = [float(row["nota_conteudo_institucional"]) for row in corr_rows]
    pearson_r = pearson(xs, ys)
    spearman_rho = pearson(rankdata(xs), rankdata(ys)) if xs else None

    lines = [
        "# Relatório automático dos resultados AMAWeb",
        "",
        f"Gerado em: {datetime.now().astimezone().isoformat(timespec='seconds')}",
        f"Entrada: `{input_path.resolve()}`",
        f"Catálogo: `{catalog_path.resolve() if catalog_path else 'não informado'}`",
        "",
        "## Cobertura da coleta",
        "",
        f"- Universo: **{len(records)} URLs**.",
        f"- Conteúdo institucional avaliado: **{counts[STATUS_CONTENT]}**.",
        f"- Bloqueio Cloudflare: **{counts[STATUS_CLOUDFLARE]}**.",
        f"- Sem resultado AMAWeb: **{counts[STATUS_MISSING]}**.",
        "",
        "Cloudflare é contabilizado como resultado da coleta e como barreira de avaliabilidade. "
        "Sua nota AMAWeb é preservada na base, mas não entra nas estatísticas do conteúdo institucional.",
        "",
        "## Estatísticas principais",
        "",
        f"- Conteúdo institucional: média **{format_number(describe(content_scores)['media'])}**, "
        f"mediana **{format_number(describe(content_scores)['mediana'])}**, n = **{len(content_scores)}**.",
        f"- Todas as páginas retornadas, inclusive Cloudflare: média **{format_number(describe(all_scores)['media'])}**, "
        f"n = **{len(all_scores)}**.",
        f"- Pares Executivo–Legislativo: **{len(pairs)} municípios**; diferença média "
        f"Executivo − Legislativo = **{format_number(describe(paired_diffs)['media'])}**.",
        f"- Transparência × acessibilidade: Pearson r = **{format_number(pearson_r, 3)}** e "
        f"Spearman ρ = **{format_number(spearman_rho, 3)}** (n = {len(xs)}).",
        "",
        "## Erros mais frequentes",
        "",
        "| Prática | Sites | Prevalência |",
        "|---|---:|---:|",
    ]
    for row in errors[:10]:
        lines.append(
            f"| `{row['pratica']}` | {row['sites']} | {format_number(row['prevalencia_percentual'], 1)}% |"
        )
    lines.extend(
        [
            "",
            "## Arquivos gerados",
            "",
            "- `base_consolidada.csv`: uma linha por URL do universo.",
            "- `falhas_por_site.csv`: uma linha por prática falha em cada portal.",
            "- `tabela_*.csv`: tabelas completas para conferência.",
            "- `tabela_*.tex`: tabelas prontas para inclusão no LaTeX.",
            "- `figuras/*.png`: imagens em 300 dpi.",
            "- `figuras/*.pdf`: figuras vetoriais para o artigo.",
        ]
    )
    if figures:
        lines.append(f"- Figuras produzidas: {', '.join(f'`{name}`' for name in figures)}.")
    if warnings:
        lines.extend(["", "## Alertas de leitura", ""])
        lines.extend(f"- {warning}" for warning in warnings)
    output.joinpath("RELATORIO_AUTOMATICO.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    args = parse_args()
    if args.top_erros <= 0:
        print("Erro: --top-erros deve ser maior que zero.", file=sys.stderr)
        return 2
    try:
        catalog = load_catalog(args.catalogo)
        source_urls = load_source_urls(args.entrada)
        evaluations, warnings, valid_json_files = load_evaluations(args.entrada, source_urls)
        attempts = load_attempts(args.entrada)
        records, duplicate_catalog = build_records(catalog, evaluations, attempts)
        if not records:
            raise RuntimeError("Nenhum registro foi encontrado.")
        args.saida.mkdir(parents=True, exist_ok=True)

        base_fields = [
            "ordem",
            "codigo_municipio",
            "municipio",
            "poder",
            "esfera",
            "unidade_gestora",
            "indice_transparencia_2025",
            "nivel_transparencia",
            "url",
            "url_retornada_amaweb",
            "associacao_url",
            "resultado_coleta",
            "nota_amaweb",
            "nota_conteudo_institucional",
            "titulo",
            "data_avaliacao",
            "tentativas",
            "erros_a",
            "erros_aa",
            "erros_aaa",
            "testes",
            "aprovados",
            "falhos",
            "manuais",
            "elementos_html",
            "tamanho_bytes",
            "arquivo_json",
            "arquivo_pdf",
        ]
        write_csv(args.saida / "base_consolidada.csv", records, base_fields)
        failure_rows = []
        for row in records:
            for failure in row.get("falhas_detalhadas", []):
                failure_rows.append(
                    {
                        "municipio": row.get("municipio", ""),
                        "poder": row.get("poder", ""),
                        "url": row["url"],
                        **failure,
                    }
                )
        write_csv(
            args.saida / "falhas_por_site.csv",
            failure_rows,
            ["municipio", "poder", "url", "pratica", "codigo", "elementos", "descricao"],
        )

        groups = group_table(records)
        errors = error_table(records)
        pairs = paired_table(records)
        generate_tables(
            args.saida,
            records,
            groups,
            errors,
            pairs,
            valid_json_files,
            len(warnings),
            duplicate_catalog,
        )
        figures = [] if args.sem_graficos else generate_figures(
            args.saida, records, groups, errors, pairs, args.top_erros
        )
        generate_report(
            args.saida,
            records,
            groups,
            errors,
            pairs,
            warnings,
            figures,
            args.catalogo,
            args.entrada,
        )
    except RuntimeError as exc:
        print(f"Erro: {exc}", file=sys.stderr)
        return 1

    counts = Counter(row["resultado_coleta"] for row in records)
    print(f"Análise atualizada em: {args.saida.resolve()}")
    print(
        f"URLs: {len(records)} | conteúdo: {counts[STATUS_CONTENT]} | "
        f"Cloudflare: {counts[STATUS_CLOUDFLARE]} | sem resultado: {counts[STATUS_MISSING]}"
    )
    print(f"JSON válidos lidos: {valid_json_files} | URLs consolidadas com JSON: {len(evaluations)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
